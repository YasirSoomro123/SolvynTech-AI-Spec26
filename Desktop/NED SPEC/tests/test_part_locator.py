"""Part discovery: the app is given a number and a folder, nothing else.

Files may be named 2280.igs, 2280_drawing.pdf, part-2280.iges, '2280 rev B.pdf', or not
mention the number at all and only carry it inside the workbook.  The two things that
must never happen are: picking a file that belongs to another part, and crashing.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from src.part_locator import discover_parts, iter_files, locate, norm_part, part_tokens


# ---------------------------------------------------------------- naming styles
@pytest.mark.parametrize("filename,kind", [
    ("2280.igs", "IGES"), ("2280_drawing.pdf", "PDF"), ("part-2280.iges", "IGES"),
    ("2280 rev B.pdf", "PDF"), ("PROD_2280.xlsx", "Excel"), ("2280.csv", "Excel"),
])
def test_the_number_is_found_in_every_naming_style(tmp_path, filename, kind):
    (tmp_path / filename).write_bytes(b"x")
    files = locate("2280", tmp_path)
    assert kind in files.chosen, files.rejected
    assert files.chosen[kind].path.name == filename


def test_a_folder_named_after_the_part_also_counts(tmp_path):
    sub = tmp_path / "2280"
    sub.mkdir()
    (sub / "drawing.pdf").write_bytes(b"x")
    assert locate("2280", tmp_path).chosen.get("PDF")


def test_a_leading_zero_number_still_matches(tmp_path):
    (tmp_path / "002280.igs").write_bytes(b"x")
    assert locate("2280", tmp_path).chosen.get("IGES")


def test_a_partial_number_is_not_an_identity(tmp_path):
    """'12280' is not part '2280': a substring inside a longer number must not qualify."""
    (tmp_path / "12280.igs").write_bytes(b"x")
    assert locate("2280", tmp_path).chosen.get("IGES") is None


def test_a_number_inside_a_version_string_is_not_an_identity(tmp_path):
    (tmp_path / "report-v22801.pdf").write_bytes(b"x")
    assert locate("2280", tmp_path).chosen.get("PDF") is None


# ---------------------------------------------------------------- content matching
def test_a_file_is_matched_by_its_content_when_the_name_says_nothing(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    for old in list(wb.worksheets):
        wb.remove(old)
    ws = wb.create_sheet("records")
    ws.append(["part number", "material"])
    ws.append([4412, "POM"])
    wb.save(tmp_path / "master_list.xlsx")
    files = locate("4412", tmp_path)
    assert files.chosen.get("Excel"), files.rejected


def test_prose_about_another_part_does_not_claim_it(tmp_path):
    """A remark cell that merely mentions a part number is not that part's record."""
    from openpyxl import Workbook
    wb = Workbook()
    for old in list(wb.worksheets):
        wb.remove(old)
    ws = wb.create_sheet("records")
    ws.append(["part number", "remarks"])
    ws.append([1000, "See the notes for 4412, which is made of brass"])
    wb.save(tmp_path / "other.xlsx")
    assert locate("4412", tmp_path).chosen.get("Excel") is None


# ---------------------------------------------------------------- honesty
def test_an_unknown_part_finds_nothing_and_says_so(tmp_path):
    (tmp_path / "2280.pdf").write_bytes(b"x")
    files = locate("9999", tmp_path)
    assert not files.chosen
    assert set(files.missing_kinds) == {"Excel", "PDF", "IGES"}


def test_a_missing_file_type_is_reported_not_stored_and_the_rest_still_found(tmp_path):
    (tmp_path / "2280.pdf").write_bytes(b"x")
    (tmp_path / "2280.igs").write_bytes(b"x")
    files = locate("2280", tmp_path)
    assert files.missing_kinds == ["Excel"]
    assert set(files.chosen) == {"PDF", "IGES"}


def test_the_best_candidate_is_chosen_and_the_others_are_visible(tmp_path):
    (tmp_path / "2280.pdf").write_bytes(b"x")
    (tmp_path / "2280_scan_old.pdf").write_bytes(b"yy")
    files = locate("2280", tmp_path)
    assert files.chosen["PDF"].path.name == "2280.pdf"
    assert any(c.path.name.startswith("2280") for c in files.rejected)


def test_every_candidate_carries_a_reason_why_it_was_picked(synth_dir):
    files = locate("3004", synth_dir)
    chosen = files.chosen.get("PDF")
    assert chosen and chosen.reasons


def test_a_short_numeric_query_cannot_borrow_another_parts_files(tmp_path):
    """'0' is not a part number, it is a character that occurs in every dimension table."""
    (tmp_path / "factory_master.csv").write_text(
        "part_number,material\n3001,EN24\n", encoding="utf-8")
    (tmp_path / "3004  rev C.pdf").write_bytes(b"note 0 tolerance 000")
    (tmp_path / "3004  rev C.igs").write_bytes(b"$$000")
    for query in ("0", "00", "000", "7"):
        assert not locate(query, tmp_path).chosen, query
    # the real number, padded or not, still resolves to the same files
    for query in ("3004", "03004", "003004", "30 04"):
        assert set(locate(query, tmp_path).chosen) == {"PDF", "IGES"}, query


# ---------------------------------------------------------------- helpers
def test_norm_part_folds_case_padding_and_decimal_noise():
    assert norm_part(" 002280 ") == "002280"
    assert norm_part(2280.0) == "2280"
    assert norm_part("PN-2280 rev B") == "pn2280revb"
    # zero-padded numbers reach the same tokens as the plain one
    assert "2280" in part_tokens("002280")


def test_part_tokens_split_hyphenated_identifiers():
    assert "2280" in part_tokens("PN-2280-B")


def test_iter_files_skips_hidden_and_temporary(tmp_path):
    (tmp_path / "visible.pdf").write_bytes(b"x")
    (tmp_path / "~$hidden.pdf").write_bytes(b"x")
    (tmp_path / ".cache").mkdir()
    (tmp_path / ".cache" / "cached.pdf").write_bytes(b"x")
    names = {p.name for p in iter_files(tmp_path)}
    assert names == {"visible.pdf"}


def test_discover_parts_lists_the_corpus_without_being_told(synth_dir):
    parts = discover_parts(synth_dir)
    assert {"3001", "3002", "3003", "3004", "3005"} <= set(parts)


def test_discovery_never_returns_garbage(tmp_path):
    (tmp_path / "README.md").write_text("notes", encoding="utf-8")
    (tmp_path / "2280.igs").write_bytes(b"x")
    assert discover_parts(tmp_path, read_records=False) == ["2280"]


def test_real_corpus_is_found_end_to_end(data_dir):
    """The stretch goal: any part number, in one folder that mixes everything."""
    for part in ("2280", "2281"):
        files = locate(part, data_dir)
        assert set(files.chosen) == {"Excel", "PDF", "IGES"}, files.rejected
