"""Shared fixtures.

Nothing in ``data/`` - the organizers supplied no sample files, so the repo ships no
parts and a test that needed them would break on a clean checkout.  Three kinds of
input are used on purpose:

* ``tests/fixtures/reference`` - one real example set, kept out of the demo folder so
  the smoke tests have stable bytes to read,
* a synthetic corpus generated into a temporary directory at test time, and
* files invented inside the test with ``tmp_path`` - the honesty test.  The judges
  will use unseen files, so a rule that only holds for our workbook is worthless.
"""

from __future__ import annotations

import importlib.util
import os
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

os.environ.setdefault("SPEC_SKIP_INTRO", "1")      # the intro is for people, not for AppTest
REFERENCE = REPO / "tests" / "fixtures" / "reference"


@pytest.fixture(scope="session")
def repo_root() -> Path:
    return REPO


@pytest.fixture(scope="session")
def data_dir() -> Path:
    if not REFERENCE.is_dir():
        pytest.skip("tests/fixtures/reference is not present")
    return REFERENCE


@pytest.fixture(scope="session")
def synth_dir(tmp_path_factory) -> Path:
    """The generated corpus, built on demand - it is never committed."""
    out = tmp_path_factory.mktemp("synthetic")
    spec = importlib.util.spec_from_file_location(
        "make_test_data", REPO / "scripts" / "make_test_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.build(out)
    return out


# ------------------------------------------------------------------ builders
@pytest.fixture
def make_xlsx(tmp_path):
    """Write a one-sheet workbook from rows of cells; the caller chooses the header row."""
    from openpyxl import Workbook

    def _make(rows: list[list[object]], sheet: str = "Sheet1",
              filename: str | None = None) -> Path:
        wb = Workbook()
        for old in list(wb.worksheets):    # drop the auto-created empty sheet, whatever its name
            wb.remove(old)
        ws = wb.create_sheet(sheet)        # create_sheet returns a Worksheet, never Optional
        for row in rows:
            ws.append(row)
        path = tmp_path / (filename or "book.xlsx")
        wb.save(path)
        return path

    return _make


@pytest.fixture
def make_pdf(tmp_path):
    """Write a text PDF, one string per line, so label/value pairs land on the page."""
    import pymupdf

    def _make(lines: list[str], name: str = "drawing.pdf", *,
              page: tuple[int, int] = (1189, 842)) -> Path:
        doc = pymupdf.open()
        pg = doc.new_page(width=page[0], height=page[1])
        y = 100
        for line in lines:
            pg.insert_text((80, y), line, fontsize=9)
            y += 14
        path = tmp_path / name
        doc.save(path)
        doc.close()
        return path

    return _make


@pytest.fixture
def fact():
    from src.fact import Fact

    def _make(field: str, value: object, source: str, *, unit: str | None = None,
              location: str = "test", confidence: str = "code-high") -> Fact:
        return Fact(field=field, value=value, unit=unit, source_type=source,
                    source_file=f"test_{source.lower()}", location=location,
                    confidence=confidence)

    return _make
